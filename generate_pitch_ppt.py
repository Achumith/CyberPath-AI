"""
Script to generate the CyberPath AI Idea Pitch PowerPoint presentation (.pptx)
Specifically tailored for pitching the project concept to a project guide/review committee.
"""
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def create_pitch_presentation(output_filename="CyberPath_AI_Idea_Pitch.pptx"):
    prs = Presentation()
    prs.slide_width = Inches(13.333)  # 16:9 widescreen layout
    prs.slide_height = Inches(7.5)

    # Theme Colors
    COLOR_BG = RGBColor(10, 15, 29)          # Deep Midnight Navy
    COLOR_CARD = RGBColor(16, 24, 45)        # Card Background
    COLOR_CYAN = RGBColor(0, 229, 255)       # Bright Electric Cyan
    COLOR_WHITE = RGBColor(245, 247, 250)    # Crisp White
    COLOR_MUTED = RGBColor(150, 165, 185)    # Soft Slate
    COLOR_EMERALD = RGBColor(46, 213, 115)   # Mint Emerald
    COLOR_AMBER = RGBColor(255, 171, 0)      # Amber Warning
    COLOR_RED = RGBColor(255, 71, 87)        # Soft Red

    def set_slide_background(slide):
        background = slide.background
        fill = background.fill
        fill.solid()
        fill.fore_color.rgb = COLOR_BG

    def add_header(slide, title_text, category_text="PROJECT PROPOSAL & IDEA PITCH"):
        # Category Tag
        cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(0.35))
        tf_cat = cat_box.text_frame
        tf_cat.word_wrap = True
        p_cat = tf_cat.paragraphs[0]
        p_cat.text = category_text.upper()
        p_cat.font.size = Pt(10)
        p_cat.font.bold = True
        p_cat.font.color.rgb = COLOR_CYAN
        p_cat.font.name = "Arial"

        # Main Title
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.7), Inches(11.7), Inches(0.7))
        tf_title = title_box.text_frame
        tf_title.word_wrap = True
        p_title = tf_title.paragraphs[0]
        p_title.text = title_text
        p_title.font.size = Pt(22)
        p_title.font.bold = True
        p_title.font.color.rgb = COLOR_WHITE
        p_title.font.name = "Arial"

    slide_layout = prs.slide_layouts[6] # Blank

    # ==================== SLIDE 1: Title & Pitch Hook ====================
    slide1 = prs.slides.add_slide(slide_layout)
    set_slide_background(slide1)

    shape1 = slide1.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.0), Inches(1.0), Inches(11.3), Inches(5.5)
    )
    shape1.fill.solid()
    shape1.fill.fore_color.rgb = COLOR_CARD
    shape1.line.color.rgb = COLOR_CYAN
    shape1.line.width = Pt(1.5)

    tf1 = shape1.text_frame
    tf1.word_wrap = True
    tf1.margin_left = Inches(0.8)
    tf1.margin_right = Inches(0.8)
    tf1.margin_top = Inches(0.6)

    p = tf1.paragraphs[0]
    p.text = "FINAL YEAR PROJECT PROPOSAL"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = COLOR_CYAN

    p2 = tf1.add_paragraph()
    p2.text = "CyberPath AI"
    p2.font.size = Pt(36)
    p2.font.bold = True
    p2.font.color.rgb = COLOR_WHITE
    p2.space_before = Pt(8)

    p3 = tf1.add_paragraph()
    p3.text = "AI-Powered Dynamic Cyber Attack Graph & Path Prediction System"
    p3.font.size = Pt(18)
    p3.font.color.rgb = COLOR_CYAN
    p3.space_before = Pt(4)

    p4 = tf1.add_paragraph()
    p4.text = "A proposed defensive platform that maps IT networks into intelligent attack graphs, uses Machine Learning to predict high-risk breach paths, explains why paths are dangerous, and enables defenders to test security controls in simulation before deploying them."
    p4.font.size = Pt(12)
    p4.font.color.rgb = COLOR_MUTED
    p4.space_before = Pt(14)

    p5 = tf1.add_paragraph()
    p5.text = "Focus Areas: Graph Theory • Machine Learning Risk Prediction • Explainable AI • Security Simulation"
    p5.font.size = Pt(11)
    p5.font.bold = True
    p5.font.color.rgb = COLOR_EMERALD
    p5.space_before = Pt(16)

    # ==================== SLIDE 2: The Core Problem ====================
    slide2 = prs.slides.add_slide(slide_layout)
    set_slide_background(slide2)
    add_header(slide2, "1. The Real-World Problem: Why Traditional Security Fails")

    problem_cards = [
        ("Vulnerabilities Scored in Isolation", "Traditional tools give isolated CVSS scores (e.g., 'CVE-2023-123 is 8.5'). But they ignore context:\n\n• An 8.5 flaw on an isolated offline server is low priority.\n• A 6.0 flaw on an internet gateway connecting directly to the customer database is a critical disaster!", COLOR_RED),
        ("Attackers Think in Chains & Graphs", "Defenders protect individual lists of assets; attackers exploit connected chains:\n\nInternet → Web Server (Breach)\n→ App Server (Lateral Hop)\n→ Admin PC (Credential Steal)\n→ Core Database (Exfiltration)\n\nSingle-point security cannot detect this multi-hop progression.", COLOR_AMBER),
        ("Defenders Cannot Test 'What-If'", "If a security team wants to patch 3 servers or block a port, they cannot easily test the outcome in advance:\n\n• Will this eliminate the attack path?\n• How much will our overall posture improve?\n\nThey have to guess rather than simulate mathematically.", COLOR_CYAN)
    ]

    for i, (title, desc, col) in enumerate(problem_cards):
        card = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8 + i * 3.9), Inches(1.8), Inches(3.7), Inches(4.9))
        card.fill.solid()
        card.fill.fore_color.rgb = COLOR_CARD
        card.line.color.rgb = col
        card.line.width = Pt(1)

        tf = card.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.3)
        tf.margin_right = Inches(0.3)
        tf.margin_top = Inches(0.35)

        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(15)
        p.font.bold = True
        p.font.color.rgb = col

        p2 = tf.add_paragraph()
        p2.text = desc
        p2.font.size = Pt(11)
        p2.font.color.rgb = COLOR_WHITE
        p2.space_before = Pt(12)

    # ==================== SLIDE 3: The Proposed Solution ====================
    slide3 = prs.slides.add_slide(slide_layout)
    set_slide_background(slide3)
    add_header(slide3, "2. Proposed Solution: How CyberPath AI Solves This")

    sol_cards = [
        ("1. Graph-Based Network Modeling", "Model infrastructure as a directed graph:\n\n• Nodes = Computers, Web Servers, Cloud Vaults, Domain Controllers.\n• Edges = Network connectivity, permissions, and protocols.\n\nTurns complex IT environments into a clear, searchable map.", COLOR_CYAN),
        ("2. Machine Learning Risk Prediction", "Automatically discover all possible paths an attacker could take.\n\nThen, use Machine Learning (Random Forest & Gradient Boosting) to analyze path features and predict which paths are highest risk (0–100 score).", COLOR_EMERALD),
        ("3. Explainability & What-If Simulation", "• Explainable AI (XAI) tells analysts WHY a path is risky (+32% due to peak CVE, +24% internet exposure).\n• What-If Engine lets defenders toggle patches or firewall rules and see attack paths vanish in real-time.", COLOR_WHITE)
    ]

    for i, (title, desc, col) in enumerate(sol_cards):
        card = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8 + i * 3.9), Inches(1.8), Inches(3.7), Inches(4.9))
        card.fill.solid()
        card.fill.fore_color.rgb = COLOR_CARD
        card.line.color.rgb = col
        card.line.width = Pt(1)

        tf = card.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.3)
        tf.margin_right = Inches(0.3)
        tf.margin_top = Inches(0.35)

        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(15)
        p.font.bold = True
        p.font.color.rgb = col

        p2 = tf.add_paragraph()
        p2.text = desc
        p2.font.size = Pt(11)
        p2.font.color.rgb = COLOR_WHITE
        p2.space_before = Pt(12)

    # ==================== SLIDE 4: Proposed Core Modules ====================
    slide4 = prs.slides.add_slide(slide_layout)
    set_slide_background(slide4)
    add_header(slide4, "3. Proposed System Modules & Functionalities")

    modules = [
        ("Module 1: SOC Executive Dashboard", "Displays enterprise security posture (0–100), critical asset status, and risk distribution charts."),
        ("Module 2: Network & Asset Builder", "Allows creating simulated corporate assets (servers, databases, firewalls) and attaching synthetic CVEs."),
        ("Module 3: Interactive Attack Graph", "Visual canvas showing nodes colored by risk (Green, Yellow, Orange, Red) with asset telemetry."),
        ("Module 4: Attack Path Discovery Engine", "Graph traversal algorithms (bounded DFS) discovering candidate multi-hop breach vectors."),
        ("Module 5: ML Risk Prediction & XAI", "Predicts path risk scores and outputs percentage feature attributions (+32% CVSS, +24% Exposure)."),
        ("Module 6: Dynamic Attack Simulation", "Simulates step-by-step adversary movements mapped to MITRE ATT&CK techniques (T1190, T1068)."),
        ("Module 7: Predictive Next-Step Model", "Given current attacker position, forecasts the most probable next pivot with softmax probability."),
        ("Module 8: What-If Security Engine", "Tests hypothetical mitigations (patching, MFA, isolation) and displays before/after posture deltas."),
        ("Module 9: AI Security Analyst Assistant", "Provides structured explanations, SOC investigation questions, and defensive playbooks."),
        ("Module 10: Executive Audit Reports", "Generates comprehensive PDF and JSON reports summarizing environment risk posture.")
    ]

    for i, (mtitle, mdesc) in enumerate(modules):
        col_idx = i // 5
        row_idx = i % 5
        
        box = slide4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8 + col_idx * 5.9), Inches(1.8 + row_idx * 0.98), Inches(5.7), Inches(0.85))
        box.fill.solid()
        box.fill.fore_color.rgb = COLOR_CARD
        box.line.color.rgb = COLOR_CYAN if col_idx == 0 else COLOR_EMERALD
        box.line.width = Pt(1)

        tf = box.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.2)
        tf.margin_right = Inches(0.2)
        tf.margin_top = Inches(0.12)

        p = tf.paragraphs[0]
        p.text = mtitle
        p.font.size = Pt(12)
        p.font.bold = True
        p.font.color.rgb = COLOR_CYAN if col_idx == 0 else COLOR_EMERALD

        p2 = tf.add_paragraph()
        p2.text = mdesc
        p2.font.size = Pt(10)
        p2.font.color.rgb = COLOR_WHITE

    # ==================== SLIDE 5: Technical Architecture ====================
    slide5 = prs.slides.add_slide(slide_layout)
    set_slide_background(slide5)
    add_header(slide5, "4. Proposed Technical Architecture & Workflow")

    tech_steps = [
        ("1. Graph Construction", "NetworkX DiGraph modeling assets as nodes and network conduits as weighted edges with access conditions.", COLOR_CYAN),
        ("2. Feature Extraction", "Extracts 15 key features per path: path length, hops, max CVSS, avg CVSS, privilege jump, external exposure.", COLOR_EMERALD),
        ("3. Machine Learning", "Trains Random Forest & Gradient Boosting models on synthetic paths to evaluate risk with continuous precision.", COLOR_WHITE),
        ("4. Modern Full-Stack UI", "FastAPI backend + React 19 / TypeScript frontend with React Flow interactive canvas and Tailwind styling.", COLOR_CYAN)
    ]

    for i, (title, desc, col) in enumerate(tech_steps):
        card = slide5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8 + i * 2.95), Inches(1.8), Inches(2.75), Inches(4.8))
        card.fill.solid()
        card.fill.fore_color.rgb = COLOR_CARD
        card.line.color.rgb = col
        card.line.width = Pt(1)

        tf = card.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.2)
        tf.margin_right = Inches(0.2)
        tf.margin_top = Inches(0.3)

        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(14)
        p.font.bold = True
        p.font.color.rgb = col

        p2 = tf.add_paragraph()
        p2.text = desc
        p2.font.size = Pt(11)
        p2.font.color.rgb = COLOR_WHITE
        p2.space_before = Pt(12)

    # ==================== SLIDE 6: What Makes This Unique ====================
    slide6 = prs.slides.add_slide(slide_layout)
    set_slide_background(slide6)
    add_header(slide6, "5. What Makes This Unique (Not Just Another CRUD App)")

    novelty_cards = [
        ("Beyond CRUD Operations", "This is not a simple inventory manager. The core intelligence lies in mathematical graph traversal, ML risk inference, and lateral path computation.", COLOR_CYAN),
        ("Explainable AI (XAI)", "Instead of outputting an unhelpful black-box number, the system tells the analyst exactly which factors drove the score (e.g. +32% Peak CVSS, +24% Direct Ingress).", COLOR_EMERALD),
        ("Interactive 'What-If' Sandbox", "Allows defenders to pre-validate security mitigations: 'If we patch Log4j and enforce MFA on the DB, our high-risk paths drop from 8 to 2 and posture jumps by +40 points.'", COLOR_WHITE)
    ]

    for i, (title, desc, col) in enumerate(novelty_cards):
        card = slide6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8 + i * 3.9), Inches(1.8), Inches(3.7), Inches(4.9))
        card.fill.solid()
        card.fill.fore_color.rgb = COLOR_CARD
        card.line.color.rgb = col
        card.line.width = Pt(1)

        tf = card.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.3)
        tf.margin_right = Inches(0.3)
        tf.margin_top = Inches(0.35)

        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(15)
        p.font.bold = True
        p.font.color.rgb = col

        p2 = tf.add_paragraph()
        p2.text = desc
        p2.font.size = Pt(11)
        p2.font.color.rgb = COLOR_WHITE
        p2.space_before = Pt(12)

    # ==================== SLIDE 7: Future Scope & Extensions ====================
    slide7 = prs.slides.add_slide(slide_layout)
    set_slide_background(slide7)
    add_header(slide7, "6. Future Development Roadmap & Research Potential")

    future_points = [
        ("1. Autonomous Defense via Reinforcement Learning", "Train a Reinforcement Learning agent (DQN/PPO) that automatically discovers optimal defensive control placements with minimal cost.", COLOR_CYAN),
        ("2. Live Threat Intelligence Integration (STIX/TAXII)", "Ingest real-world threat feeds (CISA Known Exploited Vulnerabilities - KEV) to dynamically update CVE weights based on active exploits in the wild.", COLOR_EMERALD),
        ("3. Cloud & Kubernetes Infrastructure Ingestion", "Build automated topology connectors that parse AWS VPCs, Azure ARM, and Kubernetes cluster manifests directly into attack graphs.", COLOR_WHITE),
        ("4. Zero-Trust Identity & Access Verification", "Incorporate identity trust scores (conditional access, MFA token lifecycles) directly into graph traversal edge constraints.", COLOR_CYAN),
        ("5. Multi-Agent Red Team vs. Blue Team War Gaming", "Simulate competitive game-theoretic battles where automated red agents attempt data exfiltration while blue agents dynamically deploy honeypots.", COLOR_EMERALD)
    ]

    for i, (title, desc, col) in enumerate(future_points):
        box = slide7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.7 + i * 1.0), Inches(11.7), Inches(0.85))
        box.fill.solid()
        box.fill.fore_color.rgb = COLOR_CARD
        box.line.color.rgb = col
        box.line.width = Pt(1)

        tf = box.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.3)
        tf.margin_right = Inches(0.3)
        tf.margin_top = Inches(0.12)

        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(13)
        p.font.bold = True
        p.font.color.rgb = col

        p2 = tf.add_paragraph()
        p2.text = desc
        p2.font.size = Pt(11)
        p2.font.color.rgb = COLOR_WHITE

    # ==================== SLIDE 8: Expected Deliverables & Feasibility ====================
    slide8 = prs.slides.add_slide(slide_layout)
    set_slide_background(slide8)
    add_header(slide8, "7. Feasibility & Expected Project Outcomes")

    outcomes = [
        ("Fully Functional Web App", "Complete working SOC dashboard, interactive attack graph canvas, and dynamic simulation player.", COLOR_CYAN),
        ("Trained ML Risk Pipeline", "Reproducible dataset generator and trained Random Forest models evaluating multi-hop risk accurately.", COLOR_EMERALD),
        ("Defensive Pre-Validation", "Working What-If simulator showing mathematical before/after posture improvements and blocked paths.", COLOR_WHITE),
        ("Ethical & Self-Contained", "100% simulated defensive research. Runs smoothly locally with zero risk of real-world exploitation.", COLOR_CYAN)
    ]

    for i, (title, desc, col) in enumerate(outcomes):
        card = slide8.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8 + i * 2.95), Inches(1.8), Inches(2.75), Inches(4.8))
        card.fill.solid()
        card.fill.fore_color.rgb = COLOR_CARD
        card.line.color.rgb = col
        card.line.width = Pt(1)

        tf = card.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.2)
        tf.margin_right = Inches(0.2)
        tf.margin_top = Inches(0.3)

        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(14)
        p.font.bold = True
        p.font.color.rgb = col

        p2 = tf.add_paragraph()
        p2.text = desc
        p2.font.size = Pt(11)
        p2.font.color.rgb = COLOR_WHITE
        p2.space_before = Pt(12)

    # ==================== SLIDE 9: Pitch Conclusion ====================
    slide9 = prs.slides.add_slide(slide_layout)
    set_slide_background(slide9)

    shape9 = slide9.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.0), Inches(1.0), Inches(11.3), Inches(5.5)
    )
    shape9.fill.solid()
    shape9.fill.fore_color.rgb = COLOR_CARD
    shape9.line.color.rgb = COLOR_CYAN
    shape9.line.width = Pt(1.5)

    tf9 = shape9.text_frame
    tf9.word_wrap = True
    tf9.margin_left = Inches(0.8)
    tf9.margin_right = Inches(0.8)
    tf9.margin_top = Inches(0.6)

    p = tf9.paragraphs[0]
    p.text = "PROPOSAL SUMMARY"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = COLOR_CYAN

    p2 = tf9.add_paragraph()
    p2.text = "Why CyberPath AI is an Ideal Final-Year Project"
    p2.font.size = Pt(30)
    p2.font.bold = True
    p2.font.color.rgb = COLOR_WHITE
    p2.space_before = Pt(6)

    p3 = tf9.add_paragraph()
    p3.text = "• Strong Academic & Industry Relevance: Solves a critical blind spot in modern defensive cybersecurity.\n• Multi-Disciplinary: Combines Graph Algorithms (NetworkX) + Machine Learning (Scikit-Learn) + Explainable AI + Modern Web Architecture.\n• Realistic & Feasible: Built on synthetic IT environments with zero real-world exploitation risks.\n• High Research & Portfolio Value: Provides a solid foundation for publication and future development."
    p3.font.size = Pt(13)
    p3.font.color.rgb = COLOR_WHITE
    p3.space_before = Pt(14)

    p4 = tf9.add_paragraph()
    p4.text = "Seeking Your Guidance & Approval • Questions & Discussion"
    p4.font.size = Pt(16)
    p4.font.bold = True
    p4.font.color.rgb = COLOR_EMERALD
    p4.space_before = Pt(20)

    prs.save(output_filename)
    print(f"Pitch presentation saved to: {output_filename}")
    return output_filename

if __name__ == "__main__":
    create_pitch_presentation()
